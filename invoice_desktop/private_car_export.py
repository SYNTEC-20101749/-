"""使用原始财务模板导出私车公用费用明细（每页八条记录）。"""

from __future__ import annotations

from copy import copy, deepcopy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import os
import re
import shutil
import sys
import tempfile

from openpyxl import load_workbook


DEFAULT_TEMPLATE = Path(__file__).resolve().parent / "templates" / "private_car_v1.xlsx"
TEMPLATE_FILE_NAME = "新代私车公用费用统计表_模板.xlsx"


def get_private_car_template() -> Path:
    """模板固定保存在软件目录；首次从内置原件复制，以后不覆盖该副本。"""
    app_directory = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parents[1]
    template = app_directory / TEMPLATE_FILE_NAME
    if not template.is_file():
        if not DEFAULT_TEMPLATE.is_file():
            raise FileNotFoundError(f"未找到私车公用Excel模板：{DEFAULT_TEMPLATE}")
        try:
            with template.open("xb") as target, DEFAULT_TEMPLATE.open("rb") as source:
                shutil.copyfileobj(source, target)
        except FileExistsError:
            pass
    return template


def _set_text(cell, value: str) -> None:
    # 姓名、目的地、车牌均是文字，不应被Excel当作公式执行。
    cell.value = value
    cell.data_type = "s"


def _set_number(cell, value: str, *, money: bool = False) -> None:
    number = Decimal(value)
    cell.value = float(number)
    # Excel的0.###会把整数显示为“10.”；整数必须明确使用无小数点格式。
    cell.number_format = "0" if number == number.to_integral_value() else ("0.00" if money else "0.###")


def _fee_text(value: str) -> str:
    if not value:
        return ""
    number = Decimal(value)
    return format(number, ".0f" if number == number.to_integral_value() else ".2f")


def export_private_car_excel(
    rows: list[list[str]],
    *,
    license_plate: str,
    source_directory: str,
    output_directory: str,
    template_path: Path | None = None,
) -> Path:
    if not rows:
        raise ValueError("私车公用表没有可导出的记录")
    template = template_path or get_private_car_template()
    if not template.is_file():
        raise FileNotFoundError(f"未找到私车公用Excel模板：{template}")
    dates = sorted({row[0] for row in rows})
    date_part = dates[0] if len(dates) == 1 else f"{dates[0]}至{dates[-1]}"
    customer = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(source_directory).name).strip(" .") or "客户"
    customer = customer[:80].rstrip(" .")
    target = Path(output_directory) / f"新代私车公用费用统计表_{date_part}_{customer}.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = None
    temporary_path = None
    try:
        # 先复制完整模板，在这份副本中填写；模板原件始终不修改。
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as stream:
            temporary_path = Path(stream.name)
        shutil.copyfile(template, temporary_path)
        workbook = load_workbook(temporary_path)
        original = workbook.active
        sheets = [original]
        for page in range(1, (len(rows) + 7) // 8):
            sheet = workbook.copy_worksheet(original)
            # openpyxl复制工作表不会复制图片；独立复制SYNTEC标志与锚点。
            for image in original._images:
                sheet.add_image(deepcopy(image))
            sheets.append(sheet)
        for page, sheet in enumerate(sheets):
            sheet.title = f"私车公用{page + 1}"
            # 包括未填写的明细行，统一使用宋体10pt，保留其余边框/对齐格式。
            for cells in sheet.iter_rows(min_row=4, max_row=19, min_col=2, max_col=11):
                for cell in cells:
                    font = copy(cell.font)
                    font.name, font.sz = "宋体", 10
                    cell.font = font
            plate_font = copy(sheet["K2"].font)
            plate_font.name, plate_font.sz = "宋体", 12
            sheet["K2"].font = plate_font
            _set_text(sheet["K2"], license_plate.strip())
            for offset, values in enumerate(rows[page * 8:(page + 1) * 8]):
                row = 4 + offset * 2
                sheet.cell(row, 2, datetime.strptime(values[0], "%Y-%m-%d").date()).number_format = "yyyy-mm-dd"
                for col in (1, 2):
                    _set_text(sheet.cell(row, col + 2), values[col])
                for col in (3, 4):
                    if values[col]:
                        sheet.cell(row, col + 2, datetime.strptime(values[col], "%H:%M").time()).number_format = "h:mm"
                for col in (5, 6, 7, 8):
                    _set_number(sheet.cell(row, col + 2), values[col], money=col == 8)
                _set_text(sheet.cell(row, 11), "过路（桥）费：" + _fee_text(values[9]))
                _set_text(sheet.cell(row + 1, 11), "停车费：" + _fee_text(values[10]))
            sheet.print_area = "B1:K26"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A5
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.scale = None
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 1
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
        # 先完整写入临时文件，再替换目标，避免失败时破坏之前的已确认文件。
        workbook.save(temporary_path)
        os.replace(str(temporary_path), str(target))
        return target
    finally:
        if workbook is not None:
            workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()