from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook

from python_recognizer.types import InvoiceOrganizeResult
from python_recognizer.workflow import build_trip_summary_title, calculate_summary_total


ARCHIVE_HEADERS = [
    "归档时间",
    "出差标题",
    "出差开始日期",
    "出差结束日期",
    "来源目录",
    "输出目录",
    "打印文件数",
    "打印总份数",
    "大众运输",
    "出租车费用",
    "过路费",
    "住宿不含税",
    "住宿税额",
    "住宿合计",
    "出差补贴",
    "总计",
    "在途",
]


def get_default_archive_path(source_directory: Optional[str | Path] = None) -> Path:
    if source_directory:
        return Path(source_directory) / "发票打印汇总档案.xlsx"
    return Path.home() / "Desktop" / "发票打印汇总档案.xlsx"


def _get_trip_date_range(result: InvoiceOrganizeResult) -> tuple[str, str]:
    dates = [
        item.issue_date
        for item in result.summary.daily_breakdown
        if item.issue_date and item.issue_date != "未识别日期"
    ]
    if not dates:
        return ("", "")

    ordered_dates = sorted(set(dates))
    return (ordered_dates[0], ordered_dates[-1])


def _autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        max_length = max(len(value) for value in values)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 36)


def _ensure_workbook(archive_path: Path):
    if archive_path.exists():
        workbook = load_workbook(archive_path)
        worksheet = workbook.active
        if worksheet.max_row == 0:
            worksheet.append(ARCHIVE_HEADERS)
        elif worksheet.max_row >= 1:
            existing_headers = [worksheet.cell(row=1, column=index + 1).value for index in range(len(ARCHIVE_HEADERS))]
            if existing_headers != ARCHIVE_HEADERS:
                worksheet = workbook.create_sheet(title=f"归档{datetime.now().strftime('%Y%m%d%H%M%S')}")
                worksheet.append(ARCHIVE_HEADERS)
        return workbook, worksheet

    archive_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "打印归档"
    worksheet.append(ARCHIVE_HEADERS)
    return workbook, worksheet


def append_print_archive(
    result: InvoiceOrganizeResult,
    *,
    archive_path: Path | None = None,
    printed_jobs: int,
    unique_files: int,
    taxi_amount: float,
    subsidy_amount: float,
    in_transit_amount: float,
    travel_in_transit: str,
) -> Path:
    """生成当前汇总归档；同名文件存在时覆盖，不追加历史记录。"""
    target_path = archive_path or get_default_archive_path(result.source_directory)
    if target_path.exists():
        target_path.unlink()
    workbook, worksheet = _ensure_workbook(target_path)
    start_date, end_date = _get_trip_date_range(result)

    worksheet.append(
        [
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            build_trip_summary_title(result),
            start_date,
            end_date,
            result.source_directory,
            result.output_directory,
            unique_files,
            printed_jobs,
            round(result.summary.transport_total, 2),
            round(taxi_amount, 2),
            round(result.summary.toll_total, 2),
            round(result.summary.lodging_amount_total, 2),
            round(result.summary.lodging_tax_total, 2),
            round(result.summary.lodging_total, 2),
            round(subsidy_amount + in_transit_amount, 2),
            calculate_summary_total(result, subsidy_amount, taxi_amount, in_transit_amount),
            travel_in_transit,
        ]
    )

    _autosize_columns(worksheet)
    workbook.save(target_path)
    workbook.close()
    return target_path