from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "发票管理系统操作说明书-v1.0.9.docx"


def set_cell_shading(cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell_properties.append(shading)


def set_cell_text(cell, value: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], "2F6F98")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered_steps(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in [("Title", 24, "17324D"), ("Heading 1", 16, "17324D"), ("Heading 2", 13, "2F6F98")]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def add_title_page(document: Document) -> None:
    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("发票管理系统\n操作说明书")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(27)
    run.font.color.rgb = RGBColor(23, 50, 77)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("版本：v1.0.9\n适用对象：日常出差发票整理、汇总、打印及归档人员")
    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run("本说明书覆盖桌面版主要功能与标准操作流程。")
    document.add_page_break()


def build_manual() -> None:
    document = Document()
    configure_document(document)
    add_title_page(document)

    document.add_heading("目录", level=1)
    add_bullets(document, [
        "1. 软件简介与适用范围",
        "2. 启动前准备",
        "3. 标准操作流程：分析、确认、打印、归档",
        "4. 页面与功能说明",
        "5. 网约车行程单与每日打车间隔",
        "6. 汇总清单、打印与 Excel 归档",
        "7. 设置：打印机、开机自启动与报销周期闪烁",
        "8. 异常处理与常见问题",
        "9. 数据、域控部署与使用注意事项",
    ])
    document.add_page_break()

    document.add_heading("1. 软件简介与适用范围", level=1)
    document.add_paragraph(
        "发票管理系统用于将一个出差或报销周期内的 PDF 发票集中整理。系统会识别票据类别、实际发生日期、金额、税额、购方税号等信息，"
        "自动生成重命名文件、按日期汇总清单、打印队列及 Excel 归档文件。"
    )
    document.add_heading("1.1 主要能力", level=2)
    add_bullets(document, [
        "扫描指定文件夹内的 PDF；优先读取 PDF 文本层，文本不足时自动使用本地 OCR 识别。",
        "识别网约车、网约车行程单、火车票、高速通行票、住宿票等常用票据。",
        "自动匹配网约车发票与行程单，行程单不重复计入交通费用。",
        "识别高速通行费电子发票的发票号码并避免误用发票代码；高速通行费行程单将与对应发票使用同一发票号码命名。",
        "按实际发生日期汇总大众运输、出租车费用、过路费、住宿不含税金额、住宿税额、住宿合计、在途及出差补贴。",
        "自动生成汇总清单 PDF、横向 A5 的汇总合成 PDF、整理后的 PDF 副本、可打印队列和发票打印汇总档案 Excel。",
        "对购方统一社会信用代码、住宿票种、未配对行程单、金额或发票号码缺失等情况给出复核提示。",
    ])
    document.add_heading("1.2 支持的票据与计费规则", level=2)
    add_table(document, ["票据类别", "汇总处理", "打印份数", "说明"], [
        ["网约车发票", "计入大众运输", "1", "与行程单配对后，以发票金额计费。"],
        ["网约车行程单", "不计入费用", "1", "用于配对及统计每日打车间隔。"],
        ["火车票", "计入大众运输", "2", "系统按实际乘车日期汇总。"],
        ["高速通行票", "计入过路费", "1", "按实际通行日期汇总。"],
        ["住宿票", "计入住宿金额", "2", "区分专票、普票；专票可显示不含税金额和税额。"],
    ])

    document.add_heading("2. 启动前准备", level=1)
    document.add_heading("2.1 文件准备", level=2)
    add_bullets(document, [
        "将本次出差需要处理的 PDF 发票和网约车行程单放入同一个文件夹。建议一个文件夹只放一趟出差或一个报销批次的文件。",
        "不要将旧的“汇总清单.pdf”放在待处理目录中；系统会在输出目录重新生成汇总清单。",
        "网约车行程单应与对应网约车发票一并放入目录。文件名中包含编号、金额、“发票”或“行程单”时，匹配准确率更高。",
        "请保持 PDF 可正常打开。扫描件允许使用，但识别结果需在明细表中复核。",
    ])
    document.add_heading("2.2 建议的目录示例", level=2)
    document.add_paragraph("示例：D:\\出差报销\\2026-07杭州出差\\，其中放置火车票、住宿发票、网约车发票和滴滴行程单 PDF。")
    document.add_paragraph("执行分析后，系统会在该目录下创建“原文件夹名称-NewName”输出目录，并覆盖同名旧输出内容。")
    document.add_heading("3. 标准操作流程", level=1)
    document.add_paragraph("建议每次按以下顺序操作，以确保汇总、打印和归档使用的是同一批识别结果。")
    add_numbered_steps(document, [
        "启动“发票管理系统”。",
        "点击“选择文件夹”，选择存放本次 PDF 发票的目录。",
        "点击“分析汇总”。系统会识别票据、进行网约车配对、生成重命名副本与“汇总清单.pdf”。处理期间请不要关闭软件。",
        "在“发票明细”中检查类别、日期、金额、购方税号及提示。出现黄色提示或“待复核”时，请按原始 PDF 人工确认。",
        "如需补录出租车费用、在途或出差补贴，点击“填写补贴/出租车”，填写后点击“确认补贴/出租车”。",
        "点击“生成汇总清单并打开”，生成并查看 Excel 归档文件。",
        "确认无误后，点击“一键打印”打印汇总清单和全部整理后的 PDF；或点击“只打印汇总清单”。",
        "如需备份整理后的 PDF，点击“打开输出目录”；如需上传公共盘，点击“pdf上传到局域公共盘”。",
    ])

    document.add_heading("4. 页面与功能说明", level=1)
    document.add_heading("4.1 操作步骤区域", level=2)
    add_table(document, ["按钮", "用途", "操作提示"], [
        ["选择文件夹", "选择待识别 PDF 所在目录", "切换目录后应重新执行分析汇总。"],
        ["分析汇总", "识别、配对、重命名并生成汇总 PDF", "会创建或覆盖“目录名-NewName”输出目录。"],
        ["填写补贴/出租车", "显示按日期汇总表", "用于填写系统无法从票据中得出的费用。"],
        ["生成汇总清单并打开", "生成发票打印汇总档案.xlsx", "会覆盖当前来源目录内同名归档文件。"],
        ["打开输出目录", "打开整理后的 PDF 文件夹", "需先完成分析汇总。"],
        ["pdf上传到局域公共盘", "复制输出目录内的票据 PDF 到指定公共盘", "按发票号码筛选票据及配套行程单，不上传汇总文件。"],
        ["一键打印", "打印汇总清单和全部票据", "纸张为 A5；住宿票、火车票默认打印 2 份。"],
        ["手动填写汇总", "不使用票据识别结果，手工新建按日期汇总", "进入后将不使用当前分析数据。"],
        ["只打印汇总清单", "仅打印汇总清单 PDF", "适用于重新打印封面汇总表。"],
        ["出差单填写", "打开公司出差单填写入口", "需要网络及相应登录权限。"],
    ])

    document.add_heading("4.2 发票明细表", level=2)
    document.add_paragraph("发票明细表用于处理前复核。每一行对应一个原始 PDF。重点列说明如下：")
    add_table(document, ["列名", "含义", "复核要点"], [
        ["日期", "票据实际发生日期", "应与行程、入住、通行或乘车日期一致。"],
        ["类别", "系统识别的票据类型", "“待核验”需人工检查。"],
        ["总额/金额/税额", "票据金额字段", "住宿专票应重点检查金额与税额。"],
        ["供应商/平台", "销售方或出行平台", "用于辅助网约车配对与人工核对。"],
        ["发票号", "识别出的发票号码", "缺失时会显示复核提示。"],
        ["重命名文件名", "输出目录中使用的新文件名", "避免直接修改原始文件。"],
        ["配对状态", "网约车发票与行程单的匹配可信度", "未配对行程单需要人工确认。"],
        ["复核状态/提示", "系统提醒信息", "包括税号、住宿票种、金额和每日打车间隔等。"],
    ])

    document.add_heading("4.3 按日期汇总表", level=2)
    document.add_paragraph("点击“填写补贴/出租车”后显示。系统自动写入的列不可修改；“出租车费用”“在途”“出差补贴”允许填写。")
    add_table(document, ["列", "来源或填写方式"], [
        ["日期", "系统按实际发生日期归集；手动汇总模式下可自行填写。"],
        ["大众运输", "自动汇总网约车发票及火车票。"],
        ["出租车费用", "人工填写的非网约车出租车费用。"],
        ["过路费", "自动汇总高速通行票。"],
        ["住宿不含税/税额/合计", "自动汇总住宿票。"],
        ["在途", "人工填写在途金额或说明；可参与汇总。"],
        ["出差补贴", "人工填写；点击确认后写入汇总清单。"],
    ])

    document.add_heading("5. 网约车行程单与每日打车间隔", level=1)
    document.add_paragraph(
        "系统会从滴滴等网约车行程单的“上车时间”表格中读取当天各笔行程时间。对于同一天的所有行程单，"
        "系统计算最早上车时间与最晚上车时间的差值，并以小时显示。"
    )
    document.add_heading("5.1 计算规则", level=2)
    add_bullets(document, [
        "识别行程单中的时间格式，例如“07-20 07:54 周一”“07-20 18:35 周一”。",
        "同一份行程单存在多笔行程时，先取该行程单中的最早和最晚时间。",
        "同一日期存在多份行程单时，再取全部行程单中的最早和最晚时间。",
        "例如：最早上车时间为 07:54、最晚上车时间为 18:35，间隔为 10 小时 41 分钟，显示为“打车间隔：10.7h（07:54-18:35）”。",
        "每日打车间隔仅显示在“网约车行程单”对应行的“提示”栏，不显示在普通“网约车发票”行。",
        "该间隔同时写入 Excel 归档的“备注”列，便于报销复核。",
    ])
    document.add_heading("5.2 无法识别时间时", level=2)
    add_bullets(document, [
        "确认行程单 PDF 内是否含“上车时间”及具体时间；若为模糊扫描件，建议重新导出清晰 PDF。",
        "确认行程单与对应发票在同一待处理目录内。",
        "重新执行“分析汇总”后检查行程单的提示栏。",
        "若仍无法识别，可在报销资料中人工备注，系统不会因无法计算时间间隔而停止其他发票整理。",
    ])

    document.add_heading("6. 汇总清单、打印与 Excel 归档", level=1)
    document.add_heading("6.1 汇总清单 PDF", level=2)
    document.add_paragraph("分析完成后，输出目录会生成“汇总清单.pdf”。内容包括按日期汇总表、备注列和报销流程检查项。备注列会显示每日打车间隔。")
    document.add_paragraph("同时会生成“汇总合成.pdf”，该文件按实际打印顺序合并汇总清单及各票据，使用横向 A5 版面，便于预览、留存或统一打印。")
    document.add_heading("6.2 打印", level=2)
    add_bullets(document, [
        "点击“一键打印”前，先在“设定 → 打印机设定”中选择可用打印机。",
        "一键打印会先打印 1 份汇总清单，再按各票据的默认份数打印整理后的 PDF。",
        "默认打印规则：住宿票和火车票各 2 份；其他票据 1 份。",
        "打印机不可用、未设置默认打印机或纸张设置异常时，系统会显示提示。",
    ])
    document.add_heading("6.3 Excel 归档", level=2)
    document.add_paragraph("点击“生成汇总清单并打开”后，系统会在来源目录生成“发票打印汇总档案.xlsx”，若已有同名文件将覆盖。归档文件包含出差日期、来源与输出目录、打印数量、各类费用、总计和备注。")
    add_bullets(document, [
        "备注列：写入“日期：打车间隔：xh（最早-最晚）”。",
        "“住宿费”=住宿专票不含税金额+住宿普票全额；“税金”仅记录住宿专票税额。",
        "Excel 用于归档与复核，不替代原始 PDF 发票。",
    ])
    document.add_heading("6.4 公共盘上传", level=2)
    add_numbered_steps(document, [
        "完成分析汇总后，点击“pdf上传到局域公共盘”。",
        "在弹窗中确认待上传文件列表，填写或点击“选择文件夹”选择目标公共盘路径。列表会包含文件名含已识别发票号码的票据 PDF 及配套网约车行程单。",
        "点击 YES 后，系统复制票据 PDF 到目标目录；“汇总清单.pdf”和“汇总合成.pdf”不会被上传。",
        "上传成功后，系统会记住本次目标路径，下次自动带出。",
    ])

    document.add_heading("7. 设置：打印机、开机自启动与报销周期闪烁", level=1)
    document.add_heading("7.1 打开设置", level=2)
    document.add_paragraph("从顶部菜单选择“设定”，打印机、开机自启动和报销周期闪烁均通过独立设定弹窗进行配置。")
    document.add_heading("7.2 打印机设定", level=2)
    add_numbered_steps(document, [
        "从顶部菜单选择“设定 → 打印机设定”。",
        "点击“选择打印机”。",
        "从 Windows 当前可用打印机列表选择目标设备。",
        "系统会保存选择，下次启动继续使用；如设备变更，请重新选择。",
    ])
    document.add_heading("7.3 开机自启动", level=2)
    add_bullets(document, [
        "从顶部菜单选择“设定 → 开机自启动”，勾选“开启开机自启动”。",
        "开启后，当前 Windows 用户登录时将自动启动本软件。",
        "取消勾选即可关闭，不影响已生成的汇总、归档或设置。",
        "此功能使用当前用户的 Windows 启动项，无需管理员权限。",
    ])
    document.add_heading("7.4 报销周期提醒闪烁", level=2)
    add_bullets(document, [
        "从顶部菜单选择“设定 → 报销周期闪烁设定”。",
        "勾选“开启报销周期提醒闪烁”后，报销截止日期剩余 10 天或更少时，顶部提醒会闪烁。",
        "取消勾选即可关闭闪烁；红色临界提醒仍会保留。设置会自动保存。",
    ])

    document.add_heading("8. 异常处理与常见问题", level=1)
    add_table(document, ["现象", "可能原因", "处理方法"], [
        ["分析后没有文件", "选择的目录中没有 PDF", "确认目录正确，且文件扩展名为 .pdf。"],
        ["票据显示待核验", "文本层或 OCR 未能可靠识别类别", "打开原始 PDF 人工确认类别、金额与日期。"],
        ["提示购方税号未识别/不匹配", "票面税号缺失、模糊或不是指定税号", "核对购买方统一社会信用代码，必要时更换清晰 PDF。"],
        ["网约车行程单未配对", "发票与行程单金额/编号不一致或缺少其中一份", "将对应发票与行程单放入同一目录，并核对文件名、金额和日期。"],
        ["没有每日打车间隔", "未识别到行程单上车时间", "使用含上车时间表格的清晰行程单，重新分析。"],
        ["无法打印", "未检测到可用打印机", "在 Windows 中安装或连接打印机，再到设置中选择。"],
        ["公共盘上传失败", "路径不可访问或权限不足", "确认网络、共享路径与写入权限；重新选择目标文件夹。"],
        ["Excel 未更新", "未重新生成归档", "分析完成后再次点击“生成汇总清单并打开”。"],
        ["高速票文件名使用发票代码", "旧版本或旧处理结果", "重新使用更新后的程序处理；以票面“发票号码”为准。"],
    ])

    document.add_heading("9. 数据、域控部署与使用注意事项", level=1)
    add_bullets(document, [
        "系统会复制并重命名 PDF 到输出目录，不直接修改原始 PDF；原始票据仍应妥善保存。",
        "每次分析会替换同一输出目录中的旧内容。若需要保留旧结果，请在重新分析前备份输出目录。",
        "OCR 与规则识别可能受扫描质量、票面版式和文本层完整度影响；报销前必须以原始票据为准进行人工复核。",
        "购方税号、住宿票种、发票号码、金额以及网约车配对状态均属于关键复核项。",
        "请勿将非本次报销的 PDF 混入待处理目录，以免进入汇总或打印队列。",
        "域控环境部署时请使用发布包中的 SYNTEC-InvoiceManager.exe，并保留同级 _internal 文件夹；不要只单独复制 EXE。",
        "域控环境重新打包时，必须在纯英文且无空格的项目路径执行；EXE 文件名必须以 SYNTEC 开头，且禁止启用 UPX 压缩。",
        "发布包版本信息的 CompanyName 与 LegalCopyright 必须包含 SYNTEC，语言代码应使用中性配置 000004B0 / [0, 1200]；打包后确认 _internal 中存在 Python DLL 与 _ctypes 扩展。",
        "域控环境中不要以 ctypes 调用 Windows API 隐藏控制台；GUI 程序应使用 PyInstaller 的 --windowed 参数。",
    ])

    document.add_paragraph()
    closing = document.add_paragraph("文档版本：v1.0.9　　生成日期：2026-08-18")
    closing.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    closing.runs[0].font.color.rgb = RGBColor(107, 98, 87)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_manual()
