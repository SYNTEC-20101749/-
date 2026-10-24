from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import fitz
from openpyxl import load_workbook

from invoice_desktop.private_car_export import (
    DEFAULT_TEMPLATE, TEMPLATE_FILE_NAME, export_private_car_excel, get_private_car_template,
)
from invoice_desktop.printing import build_print_queue, build_summary_sheet_queue, print_pdf_queue
from invoice_desktop.spreadsheet_printing import render_excel_to_pdf


def write_pdf(path):
    with fitz.open() as doc:
        doc.new_page(width=595.276, height=419.528).insert_text((30, 30), "A5 test")
        doc.save(path)


class SpreadsheetPrintingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_template_created_beside_exe_once_and_never_overwritten(self):
        with patch("invoice_desktop.private_car_export.sys.frozen", True, create=True), \
             patch("invoice_desktop.private_car_export.sys.executable", str(self.root / "SYNTEC-InvoiceManager.exe")):
            template = get_private_car_template()
            self.assertEqual(template, self.root / TEMPLATE_FILE_NAME)
            self.assertEqual(template.read_bytes(), DEFAULT_TEMPLATE.read_bytes())
            template.write_bytes(b"user modified template")
            self.assertEqual(get_private_car_template().read_bytes(), b"user modified template")

    def test_export_copies_template_and_sets_fonts_and_a5_for_all_sheets(self):
        template = self.root / TEMPLATE_FILE_NAME
        template.write_bytes(DEFAULT_TEMPLATE.read_bytes())
        before = template.read_bytes()
        row = ["2026-10-23", "李伟", "客户甲", "9:00", "17:40", "1", "2", "1", "1.50", "", ""]
        with patch("invoice_desktop.private_car_export.shutil.copyfile", wraps=__import__("shutil").copyfile) as copy:
            result = export_private_car_excel([row] * 9, license_plate="苏U930FQ", source_directory=str(self.root / "客户甲"),
                                               output_directory=str(self.root / "renamed"), template_path=template)
        copy.assert_called_once()
        self.assertEqual(copy.call_args.args[0], template)
        self.assertEqual(result.parent, self.root / "renamed")
        self.assertEqual(template.read_bytes(), before)
        workbook = load_workbook(result)
        self.addCleanup(workbook.close)
        for sheet in workbook:
            for coordinate in ("B4", "C4", "D4", "K4", "K5", "B18", "C18", "K19"):
                self.assertEqual(sheet[coordinate].font.name, "宋体")
                self.assertEqual(sheet[coordinate].font.sz, 10)
            self.assertEqual(sheet["K2"].font.name, "宋体")
            self.assertEqual(sheet["K2"].font.sz, 12)
            self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A5))
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertIsNone(sheet.page_setup.scale)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 1)
            self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)

    def test_full_queue_adds_workbook_once_without_changing_invoice_copies(self):
        summary = self.root / "summary.pdf"
        summary.touch()
        invoice = self.root / "invoice.pdf"
        invoice.touch()
        excel = self.root / "private.xlsx"
        excel.touch()
        record = SimpleNamespace(issue_date="2026-10-23", category="住宿票", source_file="invoice.pdf", rename_target="invoice.pdf", print_copies=2)
        result = SimpleNamespace(summary_sheet_pdf=str(summary), source_directory=str(self.root),
                                 records=[record], private_car_excel=str(excel))
        self.assertEqual(build_print_queue(result), [summary, invoice, invoice, excel])
        self.assertEqual(build_summary_sheet_queue(result), [summary])
        excel.unlink()
        self.assertEqual(build_print_queue(result), [summary, invoice, invoice])
        result.private_car_excel = ""
        self.assertEqual(build_print_queue(result), [summary, invoice, invoice])

    def test_excel_integer_numbers_have_no_decimal_point_and_fractions_keep_precision(self):
        rows = [
            ["2026-10-23", "李伟", "客户甲", "9:00", "18:00", "10", "80", "70", "105.00", "20.00", "0"],
            ["2026-10-24", "李伟", "客户甲", "9:00", "18:00", "10.125", "80.375", "70.25", "105.38", "12.50", "0.50"],
        ]
        path = export_private_car_excel(rows, license_plate="苏U930FQ", source_directory=str(self.root / "客户甲"), output_directory=str(self.root))
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual([sheet.cell(4, col).value for col in range(7, 11)], [10, 80, 70, 105])
        self.assertTrue(all(sheet.cell(4, col).number_format == "0" for col in range(7, 11)))
        self.assertEqual(sheet["K4"].value, "过路（桥）费：20")
        self.assertEqual(sheet["K5"].value, "停车费：0")
        self.assertEqual(sheet["G6"].value, 10.125)
        self.assertEqual(sheet["G6"].number_format, "0.###")
        self.assertEqual(sheet["J6"].value, 105.38)
        self.assertEqual(sheet["J6"].number_format, "0.00")
        self.assertEqual(sheet["K6"].value, "过路（桥）费：12.50")
        self.assertEqual(sheet["K7"].value, "停车费：0.50")

    def test_excel_render_is_readonly_a5_and_closes_only_own_instance(self):
        excel, pdf = self.root / "private.xlsx", self.root / "private.pdf"
        excel.touch()
        workbook = MagicMock()
        sheets = [SimpleNamespace(PageSetup=SimpleNamespace()) for _ in range(2)]
        workbook.Worksheets = sheets
        workbook.ExportAsFixedFormat.side_effect = lambda **kw: write_pdf(kw["Filename"])
        application = MagicMock()
        application.Workbooks.Open.return_value = workbook
        with patch("pythoncom.CoInitialize") as initialize, patch("pythoncom.CoUninitialize") as uninitialize, \
             patch("win32com.client.DispatchEx", return_value=application) as dispatch:
            self.assertEqual(render_excel_to_pdf(excel, pdf), pdf)
        dispatch.assert_called_once_with("Excel.Application")
        initialize.assert_called_once()
        uninitialize.assert_called_once()
        self.assertEqual(application.AutomationSecurity, 3)
        application.Workbooks.Open.assert_called_once_with(Filename=str(excel.resolve()), UpdateLinks=0, ReadOnly=True)
        for sheet in sheets:
            self.assertEqual(sheet.PageSetup.PaperSize, 11)
            self.assertEqual(sheet.PageSetup.Orientation, 2)
            self.assertFalse(sheet.PageSetup.Zoom)
            self.assertEqual(sheet.PageSetup.FitToPagesWide, 1)
            self.assertEqual(sheet.PageSetup.FitToPagesTall, 1)
        workbook.Close.assert_called_once_with(SaveChanges=False)
        application.Quit.assert_called_once()

    def test_excel_render_failure_closes_instance_and_uninitializes_com(self):
        excel = self.root / "private.xlsx"
        excel.touch()
        application = MagicMock()
        application.Workbooks.Open.side_effect = RuntimeError("Open failed")
        with patch("pythoncom.CoInitialize"), patch("pythoncom.CoUninitialize") as uninitialize, \
             patch("win32com.client.DispatchEx", return_value=application):
            with self.assertRaisesRegex(RuntimeError, "Open failed"):
                render_excel_to_pdf(excel, self.root / "out.pdf")
        application.Quit.assert_called_once()
        uninitialize.assert_called_once()

    def test_office_missing_reports_error_and_does_not_send_print_job(self):
        excel = self.root / "private.xlsx"
        excel.touch()
        with patch("pythoncom.CoInitialize"), patch("pythoncom.CoUninitialize") as uninitialize, \
             patch("win32com.client.DispatchEx", side_effect=RuntimeError("not installed")), \
             patch("invoice_desktop.printing._print_prepared_pdf_queue") as native:
            with self.assertRaisesRegex(RuntimeError, "Microsoft Excel"):
                print_pdf_queue([excel])
        native.assert_not_called()
        uninitialize.assert_called_once()

    def test_mixed_queue_uses_a5_one_copy_and_original_progress_names(self):
        pdf, excel = self.root / "summary.pdf", self.root / "private.xlsx"
        pdf.touch()
        excel.touch()
        progress = MagicMock()
        generated = []
        def render(source, target):
            generated.append(target)
            write_pdf(target)
            return target
        with patch("invoice_desktop.printing.render_excel_to_pdf", side_effect=render) as convert, \
             patch("invoice_desktop.printing.QPrinterInfo") as info, \
             patch("invoice_desktop.printing.QPrinter") as printer, \
             patch("invoice_desktop.printing.QPainter") as painter, \
             patch("invoice_desktop.printing.print_pdf_file", return_value=False) as print_file:
            selected = MagicMock()
            selected.isNull.return_value = False
            selected.printerName.return_value = "SelectedPrinter"
            info.availablePrinters.return_value = [selected]
            painter.return_value.begin.return_value = True
            count = print_pdf_queue([pdf, excel], progress_callback=progress, printer_name="SelectedPrinter")
            self.assertEqual(count, 2)
            convert.assert_called_once()
            self.assertEqual(print_file.call_count, 2)
            self.assertEqual(print_file.call_args_list[0].args[0], pdf)
            self.assertEqual(print_file.call_args_list[1].args[0], generated[0])
            printer.return_value.setPageSize.assert_called_once_with(printer.A5)
            printer.return_value.setCopyCount.assert_called_once_with(1)
            self.assertEqual(progress.call_args.args, (2, 2, excel))
            painter.return_value.end.assert_called_once()
        self.assertFalse(generated[0].exists())

    def test_conversion_failure_prevents_any_partial_printing(self):
        with patch("invoice_desktop.printing.render_excel_to_pdf", side_effect=RuntimeError("转换失败")), \
             patch("invoice_desktop.printing._print_prepared_pdf_queue") as native:
            with self.assertRaises(RuntimeError):
                print_pdf_queue([self.root / "summary.pdf", self.root / "private.xlsx"])
        native.assert_not_called()


if __name__ == "__main__":
    unittest.main()