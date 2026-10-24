from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PyQt5.QtCore import QSettings, Qt, QTimer
from PyQt5.QtTest import QTest
from PyQt5.QtWidgets import QApplication, QLineEdit

from invoice_desktop.main import MainWindow
from invoice_desktop.private_car_export import export_private_car_excel


def wait_for_confirmation(window):
    deadline = time.monotonic() + 10
    while getattr(window, "confirmation_worker", None) is not None:
        if time.monotonic() >= deadline:
            raise AssertionError("费用确认未在10秒内完成")
        QTest.qWait(5)


class PrivateCarStartupTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.settings_path = str(self.directory / "settings.ini")
        settings = QSettings(self.settings_path, QSettings.IniFormat)
        with patch("invoice_desktop.main.QSettings", return_value=settings):
            self.window = MainWindow()
        self.addCleanup(self.window.close)
        choose_patch = patch("invoice_desktop.main.QFileDialog.getExistingDirectory", return_value=str(self.directory / "output"))
        self.choose_directory = choose_patch.start()
        self.addCleanup(choose_patch.stop)
        self.window.manual_summary_source_directory = str(self.directory / "手动汇总")
        self.window.manual_output_directory = str(self.directory / "output")
        self.window.show()
        self.window.toggle_date_summary()
        self.window.private_car_checkbox.setChecked(True)
        self.car = self.window.private_car_widget
        for col, value in enumerate([
            "2026-10-24", "测试使用人", "客户甲", "9:00", "17:40",
            "10", "80", "", "", "20", "10",
        ]):
            if col not in (7, 8):
                self.car.table.item(0, col).setText(value)
        QApplication.processEvents()

    def confirm(self):
        self.assertTrue(self.window.confirm_subsidy_button.isEnabled())
        QTest.mouseClick(self.window.confirm_subsidy_button, Qt.LeftButton)
        wait_for_confirmation(self.window)

    def test_startup_click_exports_excel_without_analysis_or_summary_dates(self):
        self.assertIsNone(self.window.current_result)
        self.assertEqual(self.window.date_table.rowCount(), 0)
        with patch("invoice_desktop.main.export_summary_sheet") as pdf, \
             patch("invoice_desktop.main.export_combined_print_pdf") as combined:
            self.confirm()
        pdf.assert_not_called()
        combined.assert_not_called()
        self.assertIsNone(self.window.current_result)
        self.assertEqual(self.window.date_table.rowCount(), 0)
        path = Path(self.window.standalone_private_car_excel)
        self.assertTrue(path.is_file())
        self.assertEqual(path.parent, self.directory / "output")
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["J4"].value, 105)
        self.assertEqual(workbook.active["K4"].value, "过路（桥）费：20")
        self.assertTrue(self.window.open_output_button.isEnabled())

    def test_save_folder_is_selected_before_export_and_is_used(self):
        selected = self.directory / "自选保存位置 含空格"
        selected.mkdir()
        self.choose_directory.return_value = str(selected)
        with patch("invoice_desktop.main.export_private_car_excel", wraps=export_private_car_excel) as export:
            self.confirm()
        self.choose_directory.assert_called_once()
        self.assertEqual(self.choose_directory.call_args.args[2], self.window.manual_output_directory)
        self.assertEqual(export.call_args.kwargs["output_directory"], str(selected))
        self.assertEqual(Path(self.window.standalone_private_car_excel).parent, selected)
        self.assertFalse((self.directory / "output").exists())

    def test_cancel_save_folder_does_not_export_or_change_remembered_folder(self):
        previous = str(self.directory / "上次保存位置")
        self.window.settings.setValue("privateCar/outputDirectory", previous)
        self.window.settings.sync()
        self.choose_directory.return_value = ""
        with patch("invoice_desktop.main.export_private_car_excel") as export:
            self.confirm()
        self.choose_directory.assert_called_once()
        export.assert_not_called()
        self.assertEqual(self.window.settings.value("privateCar/outputDirectory"), previous)
        self.assertFalse(self.window.confirmation_in_progress)
        self.assertIsNone(self.window.confirmation_worker)
        self.assertTrue(self.window.confirm_subsidy_button.isEnabled())
        self.assertFalse((self.directory / "output").exists())

    def test_save_folder_is_remembered_after_reopening_and_prompted_each_time(self):
        first = self.directory / "第一次保存"
        first.mkdir()
        self.choose_directory.return_value = str(first)
        self.confirm()
        values = self.car.collect_rows(strict=True)[0]
        self.window.close()
        settings = QSettings(self.settings_path, QSettings.IniFormat)
        with patch("invoice_desktop.main.QSettings", return_value=settings):
            self.window = MainWindow()
        self.addCleanup(self.window.close)
        self.window.show()
        self.window.toggle_date_summary()
        self.window.private_car_checkbox.setChecked(True)
        self.car = self.window.private_car_widget
        for col, value in enumerate(values):
            if col not in (7, 8):
                self.car.table.item(0, col).setText(value)
        second = self.directory / "第二次保存"
        second.mkdir()
        self.choose_directory.return_value = str(second)
        self.confirm()
        self.assertEqual(self.choose_directory.call_count, 2)
        self.assertEqual(self.choose_directory.call_args.args[2], str(first))
        self.assertEqual(Path(self.window.standalone_private_car_excel).parent, second)
        self.assertEqual(QSettings(self.settings_path, QSettings.IniFormat).value("privateCar/outputDirectory"), str(second))

    def test_existing_summary_keeps_original_output_without_prompt(self):
        previous = str(self.directory / "独立私车保存位置")
        self.window.settings.setValue("privateCar/outputDirectory", previous)
        self.window.current_result = self.window._build_manual_summary_result([])
        self.window._fill_date_table(self.window.current_result)
        self.confirm()
        self.choose_directory.assert_not_called()
        self.assertTrue(Path(self.window.current_result.private_car_excel).is_file())
        self.assertEqual(Path(self.window.current_result.private_car_excel).parent, self.directory / "output")
        self.assertEqual(self.window.settings.value("privateCar/outputDirectory"), previous)

    def test_cancel_after_success_keeps_previous_excel_and_filled_rows(self):
        self.confirm()
        path = Path(self.window.standalone_private_car_excel)
        previous = path.read_bytes()
        rows = self.car.collect_rows(strict=True)
        self.choose_directory.return_value = ""
        self.confirm()
        self.assertEqual(self.car.collect_rows(strict=True), rows)
        self.assertEqual(self.window.standalone_private_car_excel, str(path))
        self.assertEqual(path.read_bytes(), previous)
        self.assertTrue(self.window.open_output_button.isEnabled())
        self.assertEqual(self.window.settings.value("privateCar/outputDirectory"), str(path.parent))

    def test_confirm_commits_the_still_focused_private_car_editor(self):
        self.car.table.editItem(self.car.table.item(0, 6))
        QApplication.processEvents()
        editor = self.car.table.findChild(QLineEdit)
        self.assertIsNotNone(editor)
        editor.selectAll()
        QTest.keyClicks(editor, "90")
        self.confirm()
        workbook = load_workbook(self.window.standalone_private_car_excel)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active["H4"].value, 90)
        self.assertEqual(workbook.active["J4"].value, 120)

    def test_background_export_keeps_ui_responsive_and_blocks_duplicate_confirm(self):
        release = threading.Event()
        self.addCleanup(release.set)
        heartbeat = []

        def slow_export(*args, **kwargs):
            if not release.wait(3):
                raise RuntimeError("界面事件循环被导出阻塞")
            return export_private_car_excel(*args, **kwargs)

        def tick():
            heartbeat.append(True)
            self.assertFalse(self.window.confirm_subsidy_button.isEnabled())
            self.window.confirm_subsidy_updates()
            release.set()

        with patch("invoice_desktop.main.export_private_car_excel", side_effect=slow_export) as export:
            QTimer.singleShot(50, tick)
            self.confirm()
            export.assert_called_once()
        self.assertEqual(heartbeat, [True])
        self.assertTrue(Path(self.window.standalone_private_car_excel).is_file())
        self.assertTrue(self.window.confirm_subsidy_button.isEnabled())

    def test_invalid_startup_trip_is_reported_without_creating_files(self):
        self.car.table.item(0, 6).setText("1")
        with patch("invoice_desktop.main.QMessageBox.warning") as warning:
            self.confirm()
        warning.assert_called_once()
        self.choose_directory.assert_not_called()
        self.assertFalse((self.directory / "output").exists())
        self.assertIsNone(self.window.confirmation_worker)

    def test_export_failure_restores_controls_and_allows_retry(self):
        previous = str(self.directory / "上次成功保存")
        self.window.settings.setValue("privateCar/outputDirectory", previous)
        with patch("invoice_desktop.main.export_private_car_excel", side_effect=PermissionError("Excel占用")), \
             patch("invoice_desktop.main.QMessageBox.critical") as critical:
            self.confirm()
        critical.assert_called_once()
        self.assertEqual(self.window.settings.value("privateCar/outputDirectory"), previous)
        self.assertEqual(self.window.standalone_private_car_excel, "")
        self.assertTrue(self.car.isEnabled())
        self.confirm()
        self.assertTrue(Path(self.window.standalone_private_car_excel).is_file())
        self.assertEqual(self.window.settings.value("privateCar/outputDirectory"), str(self.directory / "output"))

    def test_close_during_export_is_deferred_until_worker_finishes(self):
        release = threading.Event()
        self.addCleanup(release.set)
        closed = []

        def slow_export(*args, **kwargs):
            if not release.wait(3):
                raise RuntimeError("关闭保护测试超时")
            return export_private_car_excel(*args, **kwargs)

        def try_close():
            closed.append(self.window.close())
            release.set()

        with patch("invoice_desktop.main.export_private_car_excel", side_effect=slow_export):
            QTimer.singleShot(50, try_close)
            self.confirm()
        self.assertEqual(closed, [False])
        self.assertTrue(self.window.close())


if __name__ == "__main__":
    unittest.main()