from copy import deepcopy
from datetime import datetime, time
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PyQt5.QtCore import QSettings, Qt
from PyQt5.QtWidgets import QApplication, QMessageBox
from reportlab.platypus import Table

from invoice_desktop.archive import append_print_archive
from invoice_desktop.main import MainWindow, PRIVATE_CAR_SUMMARY_COLUMN
from invoice_desktop.private_car_export import DEFAULT_TEMPLATE, export_private_car_excel
from python_recognizer.workflow import calculate_summary_total


def trip_row(date="2026-10-23", start="44118", end="44165", toll="88", parking="10"):
    distance = float(end) - float(start)
    return [date, "李伟", "客户甲", "9:00", "17:40", start, end,
            str(distance), f"{distance * 1.5:.2f}", toll, parking]


class PrivateCarExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def export(self, rows, **kwargs):
        return export_private_car_excel(rows, license_plate="苏U930FQ",
                                        source_directory=str(self.directory / "客户甲"),
                                        output_directory=str(self.directory), **kwargs)

    def test_original_template_layout_values_and_logo_are_preserved(self):
        original = DEFAULT_TEMPLATE.read_bytes()
        path = self.export([trip_row()])
        self.assertEqual(path.name, "新代私车公用费用统计表_2026-10-23_客户甲.xlsx")
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet["K2"].value, "苏U930FQ")
        self.assertEqual(sheet["B4"].value, datetime(2026, 10, 23))
        self.assertEqual(sheet["B4"].number_format, "yyyy-mm-dd")
        self.assertEqual(sheet["C4"].value, "李伟")
        self.assertEqual(sheet["D4"].value, "客户甲")
        self.assertEqual(sheet["E4"].value, time(9, 0))
        self.assertEqual(sheet["F4"].value, time(17, 40))
        self.assertEqual(sheet["E4"].number_format, "h:mm")
        self.assertEqual([sheet.cell(4, col).value for col in range(7, 11)], [44118, 44165, 47, 70.5])
        self.assertEqual(sheet["K4"].value, "过路（桥）费：88")
        self.assertEqual(sheet["K5"].value, "停车费：10")
        self.assertIn("B4:B5", {str(value) for value in sheet.merged_cells.ranges})
        self.assertEqual(sheet["B1"].value, "新代科技私车公用费用统计表")
        self.assertEqual(sheet["B26"].value, "*CR-GF01-00-06 v1")
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(DEFAULT_TEMPLATE.read_bytes(), original)

    def test_more_than_eight_records_use_additional_template_sheets(self):
        rows = [trip_row() for _ in range(9)]
        rows[-1] = trip_row("2026-10-24", "10", "20", "1", "2")
        path = self.export(rows)
        self.assertIn("2026-10-23至2026-10-24", path.name)
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        self.assertEqual(len(workbook.worksheets), 2)
        first, second = workbook.worksheets
        self.assertEqual(first["B18"].value, datetime(2026, 10, 23))
        self.assertEqual(second["B4"].value, datetime(2026, 10, 24))
        self.assertEqual(second["J4"].value, 15)
        self.assertIsNone(second["B6"].value)
        self.assertTrue(all(len(sheet._images) == 1 for sheet in workbook))
        self.assertEqual(second["K2"].value, "苏U930FQ")

    def test_text_is_not_executed_as_formula_and_failed_replace_keeps_old_file(self):
        rows = [trip_row()]
        rows[0][2] = "=1+1"
        path = self.export(rows)
        workbook = load_workbook(path)
        self.assertEqual(workbook.active["D4"].data_type, "s")
        workbook.close()
        original = path.read_bytes()
        with patch("invoice_desktop.private_car_export.os.replace", side_effect=PermissionError("Excel占用")):
            with self.assertRaises(PermissionError):
                self.export(rows)
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(list(self.directory.glob("*.xlsx")), [path])

    def test_empty_rows_and_missing_template_are_reported(self):
        with self.assertRaises(ValueError):
            self.export([])
        with self.assertRaises(FileNotFoundError):
            self.export([trip_row()], template_path=self.directory / "missing.xlsx")


class PrivateCarIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        settings = QSettings(str(self.directory / "settings.ini"), QSettings.IniFormat)
        with patch("invoice_desktop.main.QSettings", return_value=settings):
            self.window = MainWindow()
        self.addCleanup(self.window.close)
        self.window.manual_summary_source_directory = str(self.directory / "客户甲")
        self.window.manual_output_directory = str(self.directory / "output")
        self.result = self.window._build_manual_summary_result([
            dict(issue_date="2026-10-23", transport_total=22.9, toll_total=9,
                 lodging_amount_total=0, lodging_tax_total=0, lodging_public_total=0, lodging_total=0),
            dict(issue_date="2026-10-24", transport_total=0, toll_total=5,
                 lodging_amount_total=0, lodging_tax_total=0, lodging_public_total=0, lodging_total=0),
        ])
        self.window.current_result = self.result
        self.window._fill_date_table(self.result)
        self.table = self.window.date_table
        self.car = self.window.private_car_widget
        self.window.private_car_checkbox.setChecked(True)

    def set_trip(self, values=None, row=0):
        values = values or trip_row()
        while self.car.table.rowCount() <= row:
            self.car.add_row()
        for col, value in enumerate(values):
            if col not in (7, 8):
                self.car.table.item(row, col).setText(value)

    def date_row(self, date):
        return next(row for row in range(self.window._date_data_row_count()) if self.table.item(row, 0).text() == date)

    def test_live_daily_totals_same_date_addition_and_disable_restore_invoice_tolls(self):
        original = deepcopy(self.result.summary.to_dict())
        self.set_trip()
        row = self.date_row("2026-10-23")
        self.assertEqual(self.table.horizontalHeader().visualIndex(PRIVATE_CAR_SUMMARY_COLUMN), 4)
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "70.50")
        self.assertFalse(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).flags() & Qt.ItemIsEditable)
        self.assertEqual(self.table.item(row, 4).text(), "88.00")
        self.assertEqual(self.table.item(self.table.rowCount() - 1, 2).text(), "186.40")
        self.set_trip(trip_row(start="10", end="20", toll="2", parking="0"), row=1)
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "85.50")
        self.assertEqual(self.table.item(row, 4).text(), "90.00")
        self.window._refresh_date_totals()
        self.assertEqual(self.table.item(row, 4).text(), "90.00")
        self.assertEqual(self.result.summary.to_dict(), original)
        self.window.private_car_checkbox.setChecked(False)
        self.assertEqual(self.table.item(row, 4).text(), "9.00")
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "")
        self.assertEqual(calculate_summary_total(self.result), 36.9)

    def test_only_exact_matching_date_is_updated_and_manual_values_preserved(self):
        self.table.item(0, 3).setText("12.34")
        self.table.item(0, 9).setText("20")
        self.table.item(0, 10).setText("30")
        self.set_trip(trip_row("2026-10-25"))
        self.assertEqual(self.window._date_data_row_count(), 2)
        self.assertEqual(self.window.current_result.private_car_daily_totals, {})
        self.assertEqual(self.table.item(0, 4).text(), "9.00")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "")
        self.assertEqual(self.window.date_manual_values["2026-10-23"], ("20", 12.34, 30))
        self.car.table.item(0, 0).setText("2026-10-23")
        self.assertEqual(self.window._date_data_row_count(), 2)
        self.assertEqual(self.table.item(0, 3).text(), "12.34")
        self.assertEqual(self.table.item(0, 4).text(), "88.00")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "70.50")
        self.car.table.item(0, 0).setText("2026-10-24")
        self.assertEqual(self.table.item(0, 4).text(), "9.00")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "")
        self.assertEqual(self.table.item(1, 4).text(), "88.00")
        self.assertEqual(self.table.item(1, PRIVATE_CAR_SUMMARY_COLUMN).text(), "70.50")

    def test_confirm_exports_real_template_and_updated_pdf_and_archive(self):
        self.set_trip()
        self.car.plate_edit.setText("苏U930FQ")
        with patch("python_recognizer.workflow.Table", wraps=Table) as tables, \
             patch("invoice_desktop.main.QMessageBox.critical") as critical:
            self.window.confirm_subsidy_updates()
        critical.assert_not_called()
        output = Path(self.result.output_directory)
        self.assertEqual(len(list(output.glob("新代私车公用费用统计表_2026-10-23_客户甲.xlsx"))), 1)
        self.assertTrue((output / "汇总清单.pdf").is_file())
        self.assertTrue((output / "汇总合成.pdf").is_file())
        data = next(call.args[0] for call in tables.call_args_list if len(call.args[0][0]) == 10)
        plain = [[cell.getPlainText() for cell in row] for row in data]
        self.assertEqual(plain[0][3], "私车公用补贴")
        self.assertEqual(plain[1][3:5], ["70.50", "88.00"])
        self.assertEqual(plain[-1][1], "186.40")
        archive = append_print_archive(self.result, archive_path=self.directory / "archive.xlsx",
                                       printed_jobs=1, unique_files=1, taxi_amount=0,
                                       subsidy_amount=0, in_transit_amount=0, travel_in_transit="")
        workbook = load_workbook(archive)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.cell(2, 11).value, 93)
        self.assertEqual(workbook.active.cell(2, 15).value, 186.4)
        self.assertEqual(workbook.active.cell(2, 18).value, 70.5)

    def test_reversed_private_rows_match_summary_dates_not_row_numbers(self):
        self.set_trip(trip_row("2026-10-24", start="0", end="10", toll="3"))
        self.set_trip(trip_row("2026-10-23", start="0", end="20", toll="4"), row=1)
        self.set_trip(trip_row("2026-10-24", start="10", end="12", toll="2"), row=2)
        self.assertEqual(self.table.item(0, 0).text(), "2026-10-23")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "30.00")
        self.assertEqual(self.table.item(0, 4).text(), "4.00")
        self.assertEqual(self.table.item(1, 0).text(), "2026-10-24")
        self.assertEqual(self.table.item(1, PRIVATE_CAR_SUMMARY_COLUMN).text(), "18.00")
        self.assertEqual(self.table.item(1, 4).text(), "5.00")

    def test_manual_summary_date_edit_rechecks_matching_and_restores_toll(self):
        self.window.manual_summary_mode = True
        self.window._fill_manual_date_table()
        self.table.item(0, 0).setText("2026-10-23")
        self.table.item(0, 4).setText("6")
        self.set_trip()
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "70.50")
        self.assertEqual(self.table.item(0, 4).text(), "88.00")
        self.table.item(0, 0).setText("2026-10-25")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "")
        self.assertEqual(self.table.item(0, 4).text(), "6")
        self.assertEqual(self.result.private_car_daily_totals, {})
        self.table.item(0, 0).setText("2026-10-23")
        self.assertEqual(self.table.item(0, PRIVATE_CAR_SUMMARY_COLUMN).text(), "70.50")
        self.assertEqual(self.table.item(0, 4).text(), "88.00")

    def test_invalid_fields_block_confirm_without_dropping_other_valid_live_amounts(self):
        self.set_trip()
        for col, value in ((0, "bad-date"), (6, "1"), (9, "NaN"), (10, "-5")):
            with self.subTest(col=col):
                previous = self.car.table.item(0, col).text()
                self.car.table.item(0, col).setText(value)
                expected = {
                    0: {}, 6: {"2026-10-23": (0.0, 88.0)},
                    9: {"2026-10-23": (70.5, 0.0)}, 10: {"2026-10-23": (70.5, 88.0)},
                }
                self.assertEqual(self.car.collect_daily_totals(), expected[col])
                with patch("invoice_desktop.main.QMessageBox.warning") as warning, \
                     patch.object(self.window, "_refresh_summary_artifacts") as refresh:
                    self.window.confirm_subsidy_updates()
                warning.assert_called_once()
                refresh.assert_not_called()
                self.car.table.item(0, col).setText(previous)

    def test_manual_mode_repeated_confirm_and_duplicate_dates_do_not_double_count(self):
        self.window.manual_summary_mode = True
        self.window._fill_manual_date_table()
        self.table.item(0, 0).setText("2026-10-23")
        self.table.item(0, 4).setText("5")
        self.window.add_manual_date_row()
        self.window._commit_date_table_editor()
        self.table.item(1, 0).setText("2026-10-23")
        self.table.item(1, 4).setText("6")
        self.set_trip()
        with patch("invoice_desktop.main.QMessageBox.critical") as critical:
            self.window.confirm_subsidy_updates()
            self.window.confirm_subsidy_updates()
        critical.assert_not_called()
        self.assertEqual(calculate_summary_total(self.window.current_result), 158.5)
        self.assertEqual(self.window.current_result.summary.toll_total, 11)
        self.assertTrue(Path(self.window.current_result.private_car_excel).is_file())
        self.assertEqual(self.table.item(self.table.rowCount() - 1, 2).text(), "158.50")
        self.window.private_car_checkbox.setChecked(False)
        self.window._refresh_summary_artifacts()
        self.assertEqual(calculate_summary_total(self.window.current_result), 11)

    def test_checked_but_empty_table_does_not_export_private_excel(self):
        self.car.user_edit.setText("默认姓名")
        with patch("invoice_desktop.main.export_private_car_excel") as export:
            self.window.confirm_subsidy_updates()
        export.assert_not_called()

    def test_excel_failure_does_not_rewrite_pdfs_or_report_success(self):
        self.set_trip()
        with patch("invoice_desktop.main.export_private_car_excel", side_effect=PermissionError("Excel占用")), \
             patch.object(self.window, "_refresh_summary_artifacts") as refresh, \
             patch.object(self.window, "_flash_confirm_subsidy_button") as flash, \
             patch("invoice_desktop.main.QMessageBox.critical") as critical:
            self.window.confirm_subsidy_updates()
        critical.assert_called_once()
        refresh.assert_not_called()
        flash.assert_not_called()

    def test_unmatched_private_date_is_not_included_in_pdf_or_archive_total(self):
        self.table.item(0, 3).setText("10")
        self.table.item(0, 9).setText("20")
        self.table.item(0, 10).setText("30")
        self.set_trip(trip_row("2026-10-25", toll="4"))
        with patch("python_recognizer.workflow.Table", wraps=Table) as tables:
            self.window._refresh_summary_artifacts()
        data = next(call.args[0] for call in tables.call_args_list if len(call.args[0][0]) == 9)
        plain = {row[0].getPlainText(): [cell.getPlainText() for cell in row] for row in data[1:]}
        self.assertNotIn("2026-10-25", plain)
        self.assertEqual(plain["2026-10-23"][2], "10.00")
        self.assertEqual(plain["总计"][1], "96.90")
        self.assertEqual(self.table.item(self.table.rowCount() - 1, 2).text(), "96.90")
        self.assertEqual(self.window.current_result.private_car_daily_totals, {})
        archive = append_print_archive(self.result, archive_path=self.directory / "archive.xlsx",
                                       printed_jobs=1, unique_files=1, taxi_amount=10,
                                       subsidy_amount=30, in_transit_amount=20, travel_in_transit="20")
        workbook = load_workbook(archive)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.cell(2, 11).value, 14)
        self.assertEqual(workbook.active.cell(2, 15).value, 96.9)

    def test_matching_date_with_user_taxi_is_kept_when_private_car_is_disabled(self):
        self.set_trip(trip_row("2026-10-23"))
        row = self.date_row("2026-10-23")
        self.table.item(row, 3).setText("10")
        self.window.private_car_checkbox.setChecked(False)
        row = self.date_row("2026-10-23")
        self.assertEqual(self.table.item(row, 3).text(), "10")
        self.assertEqual(self.table.item(row, 4).text(), "9.00")
        with patch("python_recognizer.workflow.Table", wraps=Table) as tables:
            self.window._refresh_summary_artifacts()
        data = next(call.args[0] for call in tables.call_args_list if len(call.args[0][0]) == 9)
        plain = {row[0].getPlainText(): [cell.getPlainText() for cell in row] for row in data[1:]}
        self.assertEqual(plain["2026-10-23"][2], "10.00")
        self.assertEqual(plain["总计"][1], "46.90")

    def test_one_click_print_includes_current_excel_once_and_summary_only_excludes_it(self):
        self.set_trip()
        self.window.confirm_subsidy_updates()
        self.window.selected_printer_name = "TestPrinter"
        with patch.object(self.window, "_available_printer_names", return_value=["TestPrinter"]), \
             patch("invoice_desktop.main.QMessageBox.question", return_value=QMessageBox.Yes) as question, \
             patch.object(self.window, "_start_print_worker") as start:
            self.window.print_all_documents()
            start.assert_called_once()
            queue = start.call_args.kwargs["queue"]
            excel_files = [path for path in queue if path.suffix.lower() == ".xlsx"]
            self.assertEqual(len(excel_files), 1)
            self.assertEqual(excel_files[0].parent, Path(self.result.output_directory))
            self.assertEqual(str(excel_files[0]), self.window.current_result.private_car_excel)
            self.assertIn("1份", question.call_args.args[2])
            self.assertIn("A5", question.call_args.args[2])
            start.reset_mock()
            self.window.print_summary_sheet_only()
            self.assertTrue(all(path.suffix.lower() == ".pdf" for path in start.call_args.kwargs["queue"]))

    def test_one_click_print_skips_excel_after_edit_until_confirmed_again(self):
        self.set_trip()
        self.window.confirm_subsidy_updates()
        old_excel = Path(self.window.current_result.private_car_excel)
        previous = old_excel.read_bytes()
        self.car.table.item(0, 6).setText("44166")
        self.window.selected_printer_name = "TestPrinter"
        with patch.object(self.window, "_available_printer_names", return_value=["TestPrinter"]), \
             patch("invoice_desktop.main.QMessageBox.question", return_value=QMessageBox.Yes), \
             patch.object(self.window, "_start_print_worker") as start, \
             patch("invoice_desktop.main.export_private_car_excel", wraps=export_private_car_excel) as export:
            self.window.print_all_documents()
            self.assertTrue(all(path.suffix.lower() == ".pdf" for path in start.call_args.kwargs["queue"]))
            export.assert_not_called()
            self.assertEqual(old_excel.read_bytes(), previous)
            self.window.confirm_subsidy_updates()
            export.assert_called_once()
            export.reset_mock()
            start.reset_mock()
            self.window.print_all_documents()
            excel_path = next(path for path in start.call_args.kwargs["queue"] if path.suffix.lower() == ".xlsx")
            export.assert_not_called()
            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active["J4"].value, 72)
            workbook.close()
            start.reset_mock()
            self.car.table.item(0, 9).setText("-1")
            with patch("invoice_desktop.main.QMessageBox.critical") as error:
                self.window.print_all_documents()
            error.assert_not_called()
            self.assertTrue(all(path.suffix.lower() == ".pdf" for path in start.call_args.kwargs["queue"]))

    def test_unchecked_confirm_skips_all_private_excel_processing(self):
        self.window.private_car_checkbox.setChecked(False)
        # 隐藏表即使留有无效值，也不校验、不复制模板、不导出Excel。
        self.car.table.item(0, 0).setText("invalid")
        with patch.object(self.window, "_export_current_private_car_excel") as helper, \
             patch("invoice_desktop.main.export_private_car_excel") as export, \
             patch.object(self.car, "collect_rows", side_effect=AssertionError("未勾选不应读取私车明细")), \
             patch("invoice_desktop.main.QMessageBox.warning") as warning, \
             patch("invoice_desktop.main.QMessageBox.critical") as error:
            self.window.confirm_subsidy_updates()
        helper.assert_not_called()
        export.assert_not_called()
        warning.assert_not_called()
        error.assert_not_called()
        self.assertTrue(Path(self.result.summary_sheet_pdf).is_file())

    def test_unchecking_clears_all_trip_rows_and_reference_but_keeps_remembered_settings(self):
        self.car.plate_edit.setText("苏U930FQ")
        self.car.user_edit.setText("默认使用者")
        self.set_trip()
        self.set_trip(trip_row("2026-10-25"), row=1)
        self.window.confirm_subsidy_updates()
        old_excel = Path(self.window.current_result.private_car_excel)
        self.assertTrue(old_excel.is_file())
        previous = old_excel.read_bytes()
        self.window.private_car_checkbox.setChecked(False)
        self.assertEqual(self.car.table.rowCount(), 1)
        self.assertEqual(self.car.table.item(0, 1).text(), "默认使用者")
        self.assertTrue(all(self.car.table.item(0, col).text() == "" for col in range(11) if col != 1))
        self.assertEqual(self.car.collect_rows(strict=True), [])
        self.assertEqual(self.window.current_result.private_car_excel, "")
        self.assertEqual(self.window.current_result.private_car_daily_totals, {})
        self.assertEqual(self.car.plate_edit.text(), "苏U930FQ")
        self.assertEqual(self.car.user_edit.text(), "默认使用者")
        self.assertEqual(old_excel.read_bytes(), previous)
        self.window.private_car_checkbox.setChecked(True)
        self.assertEqual(self.car.collect_rows(strict=True), [])
        self.assertEqual(self.window.current_result.private_car_excel, "")

    def test_one_click_print_without_confirm_does_not_generate_or_print_excel(self):
        self.set_trip()
        self.window.selected_printer_name = "TestPrinter"
        with patch.object(self.window, "_available_printer_names", return_value=["TestPrinter"]), \
             patch("invoice_desktop.main.QMessageBox.question", return_value=QMessageBox.Yes), \
             patch.object(self.window, "_export_current_private_car_excel") as helper, \
             patch("invoice_desktop.main.export_private_car_excel") as export, \
             patch.object(self.window, "_start_print_worker") as start:
            self.window.print_all_documents()
        helper.assert_not_called()
        export.assert_not_called()
        start.assert_called_once()
        self.assertTrue(all(path.suffix.lower() == ".pdf" for path in start.call_args.kwargs["queue"]))
        self.assertFalse(list(Path(self.result.output_directory).glob("*.xlsx")))

    def test_one_click_print_skips_deleted_excel_without_blocking_pdfs(self):
        self.set_trip()
        self.window.confirm_subsidy_updates()
        Path(self.window.current_result.private_car_excel).unlink()
        self.window.selected_printer_name = "TestPrinter"
        with patch.object(self.window, "_available_printer_names", return_value=["TestPrinter"]), \
             patch("invoice_desktop.main.QMessageBox.question", return_value=QMessageBox.Yes), \
             patch("invoice_desktop.main.QMessageBox.critical") as error, \
             patch("invoice_desktop.main.export_private_car_excel") as export, \
             patch.object(self.window, "_start_print_worker") as start:
            self.window.print_all_documents()
        error.assert_not_called()
        export.assert_not_called()
        start.assert_called_once()
        self.assertTrue(all(path.suffix.lower() == ".pdf" for path in start.call_args.kwargs["queue"]))

    def test_plate_edit_invalidates_old_excel_until_reconfirmed(self):
        self.set_trip()
        self.window.confirm_subsidy_updates()
        self.assertTrue(self.window.current_result.private_car_excel)
        self.car.plate_edit.setText("沪A12345")
        self.assertEqual(self.window.current_result.private_car_excel, "")

    def test_private_amounts_do_not_create_summary_dates_before_analysis(self):
        self.window.current_result = None
        self.table.setRowCount(0)
        self.set_trip(trip_row(start="10", end="80", toll="20"))
        self.assertEqual(self.table.rowCount(), 0)
        self.assertEqual(self.window._matched_private_car_totals(), {})
        self.assertIsNone(self.window.current_result)
        self.window.private_car_checkbox.setChecked(False)
        self.assertEqual(self.window._date_data_row_count(), 0)

    def test_confirm_button_ignores_return_and_enter_but_accepts_mouse_click(self):
        from PyQt5.QtTest import QTest
        from PyQt5.QtWidgets import QLineEdit

        self.set_trip()
        self.window._update_action_buttons()
        self.window.show()
        self.window.date_summary_group.show()
        QApplication.processEvents()
        button = self.window.confirm_subsidy_button
        self.assertTrue(button.isEnabled())
        self.assertEqual(button.text(), "确认补贴/出租车")
        self.assertTrue(button.shortcut().isEmpty())
        self.assertFalse(button.autoDefault())
        self.assertFalse(button.isDefault())
        with patch.object(self.window, "_refresh_summary_artifacts") as refresh, \
             patch("invoice_desktop.main.export_private_car_excel") as export:
            for target in (button, self.car.plate_edit, self.table):
                target.setFocus()
                QTest.keyClick(target, Qt.Key_Return)
                QTest.keyClick(target, Qt.Key_Enter)
                QApplication.processEvents()
            self.car.table.editItem(self.car.table.item(0, 9))
            QApplication.processEvents()
            editor = self.car.table.findChild(QLineEdit)
            self.assertIsNotNone(editor)
            QTest.keyClick(editor, Qt.Key_Return)
            QApplication.processEvents()
            refresh.assert_not_called()
            export.assert_not_called()
            QTest.mouseClick(button, Qt.LeftButton)
            refresh.assert_called_once()

    def test_subsidy_and_toll_sync_independently_of_time_and_parking(self):
        self.set_trip(trip_row(start="10", end="80", toll="20"))
        row = self.date_row("2026-10-23")
        self.car.table.item(0, 3).setText("未填完")
        self.car.table.item(0, 10).setText("-")
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "105.00")
        self.assertEqual(self.table.item(row, 4).text(), "20.00")
        self.car.table.item(0, 9).setText("NaN")
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "105.00")
        self.assertEqual(self.table.item(row, 4).text(), "0.00")
        self.car.table.item(0, 6).setText("")
        self.car.table.item(0, 9).setText("12.50")
        self.assertEqual(self.table.item(row, PRIVATE_CAR_SUMMARY_COLUMN).text(), "0.00")
        self.assertEqual(self.table.item(row, 4).text(), "12.50")

    def test_live_editor_updates_subsidy_and_toll_without_losing_focus(self):
        from PyQt5.QtTest import QTest
        from PyQt5.QtWidgets import QLineEdit

        self.set_trip(trip_row(start="10", end="80", toll=""))
        self.window.show()
        self.window.date_summary_group.show()
        QApplication.processEvents()
        row = self.date_row("2026-10-23")
        for col, text, expected_col, expected in ((6, "90", PRIVATE_CAR_SUMMARY_COLUMN, "120.00"), (9, "35", 4, "35.00")):
            self.car.table.editItem(self.car.table.item(0, col))
            QApplication.processEvents()
            editor = self.car.table.findChild(QLineEdit)
            self.assertIsNotNone(editor)
            editor.selectAll()
            QTest.keyClicks(editor, text)
            QApplication.processEvents()
            self.assertEqual(self.table.item(row, expected_col).text(), expected)
            self.assertTrue(editor.hasFocus())
            QTest.keyClick(editor, Qt.Key_Return)
            QApplication.processEvents()


if __name__ == "__main__":
    unittest.main()